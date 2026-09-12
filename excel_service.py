from __future__ import annotations

from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def normalize_text(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).strip()

    while "  " in text:
        text = text.replace("  ", " ")

    return text.casefold()


def normalize_date(value: Any) -> date | None:

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    formats = [
        "%d-%b-%Y",
        "%d-%B-%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d %b %Y",
        "%d %B %Y",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    return None


def _find_worksheet(workbook, worksheet_name: str):
    if worksheet_name in workbook.sheetnames:
        return workbook[worksheet_name]

    available = ", ".join(workbook.sheetnames)

    raise RuntimeError(
        f"EXCEL-007: Worksheet not found: {worksheet_name}. "
        f"Available worksheets: {available}"
    )


def _find_date_row(
    ws,
    report_date: date,
    date_column: int,
    first_data_row: int,
) -> int | None:

    for row in range(first_data_row, ws.max_row + 1):

        cell_value = ws.cell(row=row, column=date_column).value

        cell_date = normalize_date(cell_value)

        if cell_date == report_date:
            return row

    return None


def _find_ytd_row(
    ws,
    date_column: int,
) -> int | None:

    for row in range(1, ws.max_row + 1):

        value = ws.cell(row=row, column=date_column).value

        if normalize_text(value) in {
            "cum. (ytd)",
            "cum.(ytd)",
            "cum ytd",
            "cumulative",
        }:
            return row

    return None


def _find_meter_columns(
    ws,
    header_row: int,
) -> dict[str, int]:

    columns = {}

    for column in range(1, ws.max_column + 1):

        value = ws.cell(
            row=header_row,
            column=column
        ).value

        normalized = normalize_text(value)

        if normalized:
            columns[normalized] = column

    return columns


def _copy_row_format(
    ws,
    source_row: int,
    target_row: int,
) -> None:

    for column in range(1, ws.max_column + 1):

        source = ws.cell(
            row=source_row,
            column=column
        )

        target = ws.cell(
            row=target_row,
            column=column
        )

        if source.has_style:
            target._style = copy(source._style)

        if source.number_format:
            target.number_format = source.number_format

        if source.alignment:
            target.alignment = copy(source.alignment)

        if source.protection:
            target.protection = copy(source.protection)

    if source_row in ws.row_dimensions:
        ws.row_dimensions[target_row].height = (
            ws.row_dimensions[source_row].height
        )


def _create_date_row(
    ws,
    report_date: date,
    date_column: int,
    first_data_row: int,
) -> int:

    ytd_row = _find_ytd_row(
        ws,
        date_column
    )

    if ytd_row is not None:

        insert_at = ytd_row

        source_row = max(
            first_data_row,
            ytd_row - 1
        )

        ws.insert_rows(
            insert_at,
            1
        )

        _copy_row_format(
            ws,
            source_row + 1,
            insert_at
        )

        ws.cell(
            row=insert_at,
            column=date_column
        ).value = report_date

        return insert_at

    target_row = ws.max_row + 1

    _copy_row_format(
        ws,
        max(first_data_row, target_row - 1),
        target_row
    )

    ws.cell(
        row=target_row,
        column=date_column
    ).value = report_date

    return target_row


def _calculate_total(
    ws,
    row: int,
    total_column: int,
) -> float:

    total = 0.0

    for column in range(1, total_column):

        if column == 1:
            continue

        value = ws.cell(
            row=row,
            column=column
        ).value

        if isinstance(value, (int, float)):
            total += float(value)

    return round(total, 2)


def update_existing_workbook(
    workbook_path: str | Path,
    worksheet_name: str,
    report_date: date,
    meters: list[dict[str, Any]],
    header_row: int = 3,
    date_column: int = 2,
    first_data_row: int = 5,
    total_column: int = 18,
    update_total: bool = True,
    fail_on_unmapped_numeric_meter: bool = False,
) -> dict[str, Any]:

    workbook_path = Path(workbook_path).resolve()

    print(f"[EXCEL] Workbook: {workbook_path}")

    if not workbook_path.exists():
        raise RuntimeError(
            f"EXCEL-001: Workbook not found: {workbook_path}"
        )

    if not workbook_path.is_file():
        raise RuntimeError(
            f"EXCEL-002: Workbook path is not a file: {workbook_path}"
        )

    try:
        workbook = load_workbook(
            filename=workbook_path
        )
    except PermissionError as exc:
        raise RuntimeError(
            "EXCEL-003: Workbook is locked or permission was denied. "
            "Close the Excel file and try again."
        ) from exc

    try:

        ws = _find_worksheet(
            workbook,
            worksheet_name
        )

        header_columns = _find_meter_columns(
            ws,
            header_row
        )

        row = _find_date_row(
            ws,
            report_date,
            date_column,
            first_data_row
        )

        created_date_row = False

        if row is None:

            row = _create_date_row(
                ws,
                report_date,
                date_column,
                first_data_row
            )

            created_date_row = True

        updated = []
        skipped = []
        unmapped = []

        for meter in meters:

            meter_name = str(
                meter.get("meter_name", "")
            ).strip()

            active_energy = meter.get(
                "active_energy"
            )

            if not meter_name:
                skipped.append({
                    "meter": "",
                    "reason": "Missing meter name"
                })
                continue

            # N/A is NOT zero.
            if active_energy is None:
                skipped.append({
                    "meter": meter_name,
                    "reason": "Active Energy is N/A"
                })
                continue

            normalized_meter = normalize_text(
                meter_name
            )

            column = header_columns.get(
                normalized_meter
            )

            if column is None:

                unmapped.append({
                    "meter": meter_name,
                    "value": active_energy
                })

                if fail_on_unmapped_numeric_meter:
                    raise RuntimeError(
                        "EXCEL-004: Numeric PDF meter could not "
                        f"be mapped to an Excel header: {meter_name}"
                    )

                continue

            cell = ws.cell(
                row=row,
                column=column
            )

            old_value = cell.value

            cell.value = float(active_energy)

            updated.append({
                "meter": meter_name,
                "value": float(active_energy),
                "old_value": old_value,
                "cell": f"{get_column_letter(column)}{row}",
            })

        total_value = None

        if update_total:

            total_value = _calculate_total(
                ws,
                row,
                total_column
            )

            ws.cell(
                row=row,
                column=total_column
            ).value = total_value

        # Save in-place.
        workbook.save(
            workbook_path
        )

    finally:
        try:
            workbook.close()
        except Exception:
            pass

    # Re-open and verify.
    try:

        verification_wb = load_workbook(
            filename=workbook_path,
            data_only=False,
            read_only=True
        )

        verification_ws = verification_wb[
            worksheet_name
        ]

        verification_date = normalize_date(
            verification_ws.cell(
                row=row,
                column=date_column
            ).value
        )

        if verification_date != report_date:
            raise RuntimeError(
                "EXCEL-005: Post-save verification failed. "
                "Report date was not written correctly."
            )

        verification_wb.close()

    except Exception as exc:

        raise RuntimeError(
            f"EXCEL-006: Post-save verification failed: {exc}"
        ) from exc

    return {
        "success": True,
        "workbook": str(workbook_path),
        "worksheet": worksheet_name,
        "row": row,
        "report_date": report_date.isoformat(),
        "created_date_row": created_date_row,
        "updated": updated,
        "skipped": skipped,
        "unmapped": unmapped,
        "total": total_value,
    }