import shutil
from datetime import date
from pathlib import Path
import openpyxl
import pytest

from app.services.excel.service import ExcelService
from app.services.excel.mapper import ExcelMapper
from app.services.excel.verifier import ExcelVerifier


@pytest.fixture
def temp_excel(tmp_path):
    orig = Path("Test_BI_Analysis_Report_2026.xlsx")
    if not orig.exists():
        pytest.skip("Base Excel file not found")
    dest = tmp_path / "Test_Workbook.xlsx"
    shutil.copy2(orig, dest)
    return dest


def test_excel_service_headers(temp_excel):
    es = ExcelService(excel_file=temp_excel)
    headers = es.get_headers()
    assert len(headers) >= 18
    assert headers[1] == "Date"
    assert "Chiller Plating Plant" in headers
    assert "Total" in headers


def test_excel_service_existing_dates(temp_excel):
    es = ExcelService(excel_file=temp_excel)
    dates = es.get_existing_dates()
    assert len(dates) > 0
    assert date(2026, 9, 1) in dates
    assert date(2026, 9, 12) in dates


def test_excel_service_update_existing_date(temp_excel):
    es = ExcelService(excel_file=temp_excel)
    readings = [
        {"meter_name": "Chiller Plating Plant", "value": 141.35, "status": "VALID"},
        {"meter_name": "Old LT Panel Main Incomer", "value": 3165.92, "status": "VALID"},
        {"meter_name": "Packing & Stiching Line", "value": 477.82, "status": "VALID"},
        {"meter_name": "Rigga Line Fabrication", "value": 179.22, "status": "VALID"},
        {"meter_name": "Press & Fabrication", "value": 1009.43, "status": "VALID"},
        {"meter_name": "Powder Coating", "value": 2018.03, "status": "VALID"},
        {"meter_name": "New Plating Plant", "value": 1.21, "status": "VALID"},
        {"meter_name": "Old Plating Plant", "value": 2.60, "status": "VALID"},
        {"meter_name": "Kaeser ASD 60 40HP Air Compressor", "value": 776.61, "status": "VALID"},
        {"meter_name": "ELGI E18 25HP Air Compressor", "value": 0.04, "status": "VALID"},
        {"meter_name": "ELGI E45 60HP Air Compressor", "value": 1012.75, "status": "VALID"},
        {"meter_name": "DIPP_PT_PANEL", "value": 0.17, "status": "VALID"},
        {"meter_name": "DUST_ COLLECTOR", "value": 251.50, "status": "VALID"},
        {"meter_name": "SPRAY_PT_PANEL", "value": 9.19, "status": "VALID"},
        {"meter_name": "DM Plant", "value": 160.99, "status": "VALID"},
        {"meter_name": "RASKOG NEW PANL", "value": None, "status": "N/A"},
        {"meter_name": "60 HP Compressor", "value": None, "status": "N/A"},
    ]

    report_date = date(2026, 9, 12)
    res = es.update_existing_workbook(readings=readings, report_date=report_date)

    assert res.success is True
    assert res.total_value == 9206.83
    assert len(res.skipped_meters) == 2  # 2 N/A meters

    # Inspect the saved workbook
    wb = openpyxl.load_workbook(temp_excel, data_only=False)
    ws = wb["EMS Monitoring Report"]

    # Verify target row is row 16
    date_val = ws.cell(16, 2).value
    assert (hasattr(date_val, "date") and date_val.date() == report_date) or date_val == report_date

    # Verify numeric cell is float
    chiller_val = ws.cell(16, 4).value
    assert isinstance(chiller_val, (int, float))
    assert chiller_val == 141.35

    # Verify total formula is =SUM(C16:Q16)
    total_formula = ws.cell(16, 18).value
    assert total_formula == "=SUM(C16:Q16)"

    # Verify YTD formula in row 4 was preserved
    ytd_formula = ws.cell(4, 18).value
    assert ytd_formula == "=SUM(R5:R34)"

    wb.close()


def test_excel_verifier_passes(temp_excel):
    verifier = ExcelVerifier()
    # Row 15 exists in original
    res = verifier.verify(temp_excel, "EMS Monitoring Report", date(2026, 9, 11))
    assert res["valid"] is True
    assert res["target_row"] == 15


def test_excel_path_normalization_and_safety(tmp_path):
    """Verify Excel path normalization eliminates quotes, handles relative/absolute safely,
    and permanently prevents <project>\"<absolute-path> bugs."""
    base_dir = tmp_path / "app_root"
    base_dir.mkdir()

    # 1. Plain relative path
    p1 = ExcelService._normalize_path("Test_Workbook.xlsx", base_dir=base_dir)
    assert p1 == (base_dir / "Test_Workbook.xlsx").resolve()
    assert '"' not in str(p1)

    # 2. Quoted relative path
    p2 = ExcelService._normalize_path('"Test_Workbook.xlsx"', base_dir=base_dir)
    assert p2 == (base_dir / "Test_Workbook.xlsx").resolve()
    assert '"' not in str(p2)

    # 3. Single-quoted relative path with spaces
    p3 = ExcelService._normalize_path("  'Test_Workbook.xlsx'  ", base_dir=base_dir)
    assert p3 == (base_dir / "Test_Workbook.xlsx").resolve()

    # 4. Absolute path
    abs_target = (tmp_path / "actual.xlsx").resolve()
    p4 = ExcelService._normalize_path(str(abs_target), base_dir=base_dir)
    assert p4 == abs_target
    assert not str(p4).startswith(str(base_dir))

    # 5. Quoted absolute path - MUST NOT produce <base_dir>\"<abs_target>
    p5 = ExcelService._normalize_path(f'"{abs_target}"', base_dir=base_dir)
    assert p5 == abs_target
    assert '"' not in str(p5)
    assert not str(p5).startswith(str(base_dir))

    # 6. Nested quoted absolute path
    p6 = ExcelService._normalize_path(f'  " \'{abs_target}\' "  ', base_dir=base_dir)
    assert p6 == abs_target
    assert '"' not in str(p6)
    assert "'" not in str(p6)

