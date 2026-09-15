from datetime import date
from pathlib import Path
import shutil
import openpyxl
import pytest

from app.orchestration.workflow import AutomationWorkflow
from app.services.pdf.parser import PDFParser
from app.services.validation.service import ValidationService
from app.services.reconciliation.service import ReconciliationService
from app.services.excel.service import ExcelService
from app.persistence.database import Database
from app.persistence.repositories.report_repository import ReportRepository


SAMPLE_PDF = Path("data/pdfs/report_2026-09-11_2026-09-12.pdf").resolve()
ORIGINAL_EXCEL = Path("Test_BI_Analysis_Report_2026.xlsx").resolve()


class MockGmailRepository:
    def __init__(self, pdf_path: Path):
        self.pdf_path = pdf_path

    def find_reports(self):
        return [
            {
                "id": "msg_001",
                "subject": "NBSense EMS Monitoring Report - 2026-09-12",
                "file_path": str(self.pdf_path),
            }
        ]


def test_workflow_idempotency_same_report_twice(tmp_path):
    """Running the exact same report twice must update row 16 in place without row duplication."""
    test_excel = tmp_path / "Test_Workbook_Idempotent.xlsx"
    shutil.copy2(ORIGINAL_EXCEL, test_excel)

    test_db_path = tmp_path / "test_idempotent.db"
    db = Database(str(test_db_path))
    db.connect()
    db.initialize()
    repo = ReportRepository(db)

    pdf_parser = PDFParser()
    validation_service = ValidationService()
    excel_service = ExcelService(
        file_path=str(test_excel),
        sheet_name="EMS Monitoring Report",
    )
    reconciliation_service = ReconciliationService(excel_service=excel_service)
    gmail_repo = MockGmailRepository(SAMPLE_PDF)

    workflow = AutomationWorkflow(
        gmail_service=None,
        gmail_repository=gmail_repo,
        pdf_service=pdf_parser,
        validation_service=validation_service,
        reconciliation_service=reconciliation_service,
        excel_service=excel_service,
        energy_analyzer=None,
        ai_decision_engine=None,
        powerbi_service=None,
        powerbi_publisher=None,
        powerbi_refresh=None,
        notification_service=None,
        repository=repo,
    )

    # Initial sheet row count
    wb_orig = openpyxl.load_workbook(test_excel)
    orig_max_row = wb_orig["EMS Monitoring Report"].max_row
    wb_orig.close()

    # First run
    res1 = workflow.run()
    assert res1["success"] is True
    excel_res1 = res1["processed"][0]["excel"]
    assert excel_res1.success is True

    # Second run (exact duplicate)
    res2 = workflow.run()
    assert res2["success"] is True
    excel_res2 = res2["processed"][0]["excel"]
    assert excel_res2.success is True

    # Verify Excel state after duplicate run
    wb = openpyxl.load_workbook(test_excel, data_only=False)
    sheet = wb["EMS Monitoring Report"]

    # Total max rows must not have expanded
    assert sheet.max_row == orig_max_row

    # Row 16 date and formula
    assert str(sheet.cell(row=16, column=2).value).startswith("2026-09-12")
    assert sheet.cell(row=16, column=18).value == "=SUM(C16:Q16)"

    # Row 17 (next day 2026-09-13) should not have been overwritten or corrupted
    assert str(sheet.cell(row=17, column=2).value).startswith("2026-09-13")

    # Sum of row 16 meter values
    meter_values = [
        sheet.cell(row=16, column=col).value
        for col in range(3, 18)
        if sheet.cell(row=16, column=col).value is not None
    ]
    assert len(meter_values) == 15
    assert pytest.approx(sum(meter_values), rel=1e-4) == 9206.83

    wb.close()
